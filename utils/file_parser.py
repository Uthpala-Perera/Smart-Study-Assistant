import PyPDF2
from pptx import Presentation

def parse_pdf(file):
    reader = PyPDF2.PdfReader(file)
    text = ""
    for page in reader.pages:
        if page.extract_text():
            text += page.extract_text()
    return text

def parse_ppt(file):
    prs = Presentation(file)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text
    return text